"""Reading a contact list out of the file an operator actually has.

The panel takes a pasted list, and that is what the campaign is created from
-- one validated path, one place where a number becomes an MSISDN. But nobody
keeps farmer numbers as pasted text: they keep them in a spreadsheet a field
officer filled in, and asking someone to open it, select a column and paste is
how numbers get lost a row at a time.

So a file becomes those same lines, here, in Python. A spreadsheet is read
with openpyxl in this process -- the numbers do not go to a converter, a
service or the browser's memory -- and the result is handed back as
``name, number`` text for the operator to look at *before* anything is
created. Nothing is imported by this module; it turns a file into a proposal.

Three things a real spreadsheet does that a naive reader gets wrong, all of
them found in the files people send:

**A phone number stored as a number.** Excel writes 9876543210 as a float, so
``str(cell.value)`` yields "9876543210.0" and the last digit is silently a
zero. Integral floats are rendered as integers here.

**A leading zero eaten by the spreadsheet.** "09876543210" is stored as
9876543210. That is fine -- the parser accepts both -- but "+91 98765 43210"
kept as text must survive too, so text cells are never coerced.

**The header row.** Dropped by not matching, never by position: a file whose
first row is data would otherwise lose a farmer, and a file with two header
rows would keep one.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from uaagro_domain.errors import ValidationError

#: Enough for the largest campaign the compliance gate allows, several times
#: over, and small enough that a stray 200 MB export is refused rather than
#: read into memory.
MAX_IMPORT_BYTES = 8 * 1024 * 1024

#: A spreadsheet with more rows than this is not a campaign list.
MAX_ROWS = 50_000

#: What the file extension has to be. `.xls` is the pre-2007 binary format,
#: which openpyxl cannot read; it is named so the operator is told what to do
#: rather than being told the file is broken.
XLSX = ".xlsx"
LEGACY_EXCEL = (".xls",)
TEXT_SUFFIXES = (".csv", ".txt", ".tsv")

_DIGITS = re.compile(r"\D")


@dataclass(slots=True)
class Extracted:
    """What one file offered up."""

    #: ``name, number`` or bare ``number``, in file order, ready for the box
    #: the operator is looking at.
    lines: list[str] = field(default_factory=list)
    #: Rows that carried no number: header rows, blank rows, notes, totals.
    skipped: int = 0
    #: How many sheets were read, for a workbook with more than one.
    sheets: int = 1


def looks_like_number(cell: str) -> bool:
    """Whether a cell reads as an Indian mobile number.

    Deliberately loose: this decides which cell in a row is the number, not
    whether the number is callable. ``normalise_msisdn`` decides that later,
    and a cell that gets this far and fails there is reported to the operator
    as an unreadable line rather than dropped.
    """
    digits = _DIGITS.sub("", cell)
    if not 10 <= len(digits) <= 13:
        return False
    # Strip the country code and a trunk zero, then India's mobile range.
    trimmed = digits[2:] if len(digits) >= 12 and digits.startswith("91") else digits
    trimmed = trimmed.lstrip("0")
    return len(trimmed) == 10 and trimmed[0] in "6789"


def _cell_text(value: Any) -> str:
    """One spreadsheet cell as the text it was meant to be.

    An integral float is a phone number Excel decided was a quantity; anything
    else is left as its own string representation.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def _line_from(cells: list[str]) -> str | None:
    """``name, number`` from one row, or None when there is no number."""
    filled = [cell for cell in (c.strip() for c in cells) if cell]
    number = next((cell for cell in filled if looks_like_number(cell)), None)
    if number is None:
        return None
    name = next((cell for cell in filled if cell != number and not looks_like_number(cell)), None)
    # A name with a comma in it would split the line the parser reads, so the
    # separator inside it is dropped rather than the name.
    clean = name.replace(",", " ").strip() if name else None
    clean = " ".join(clean.split()) if clean else None
    return f"{clean}, {number}" if clean else number


def extract_contacts(data: bytes, *, filename: str) -> Extracted:
    """Turn an uploaded contact list into lines for the operator to check.

    Raises:
        ValidationError: when the file cannot be read at all, naming what to
            do about it.
    """
    if len(data) > MAX_IMPORT_BYTES:
        raise ValidationError(
            f"That file is {len(data) // (1024 * 1024)} MB.",
            remedy=f"Upload a list under {MAX_IMPORT_BYTES // (1024 * 1024)} MB.",
        )
    if not data:
        raise ValidationError("The file is empty.", remedy="Choose a file with numbers in it.")

    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if suffix == XLSX:
        return _from_xlsx(data)
    if suffix in LEGACY_EXCEL:
        raise ValidationError(
            "That is the older Excel format (.xls).",
            remedy="Open it and choose Save As → Excel Workbook (.xlsx), or CSV.",
        )
    if suffix in TEXT_SUFFIXES or suffix == "":
        return _from_text(data)
    raise ValidationError(
        f"Cannot read {filename!r}.",
        remedy="Upload a spreadsheet (.xlsx) or a CSV file, or paste the numbers in.",
    )


def _from_text(data: bytes) -> Extracted:
    """CSV, tab-separated or one number per line."""
    text = data.decode("utf-8-sig", errors="replace")
    out = Extracted()
    rows = 0
    # Sniffing rather than assuming a comma: exports from Indian ERPs are
    # semicolon-separated often enough to matter, and a tab file read as CSV
    # yields one cell that happens to contain the number, which still works
    # but loses the name.
    sample = text[:4096]
    delimiter = max([",", ";", "\t", "|"], key=sample.count)
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        rows += 1
        if rows > MAX_ROWS:
            break
        line = _line_from(list(row))
        if line is None:
            out.skipped += 1
        else:
            out.lines.append(line)
    return out


def _from_xlsx(data: bytes) -> Extracted:
    """Every sheet of a workbook, values only.

    ``data_only`` returns the last value Excel calculated for a formula rather
    than the formula itself, which is what a number built with CONCATENATE
    needs. A workbook saved by a tool that never calculated returns None for
    those cells, and they are counted as skipped rather than guessed at.
    """
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise ValidationError(
            "Spreadsheets cannot be read on this server.",
            remedy="Save the list as CSV, or paste the numbers in.",
        ) from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(
            "That spreadsheet could not be opened.",
            remedy="Check it opens in Excel, or save it as CSV.",
        ) from exc

    out = Extracted(sheets=0)
    rows = 0
    try:
        for sheet in workbook.worksheets:
            out.sheets += 1
            for row in sheet.iter_rows(values_only=True):
                rows += 1
                if rows > MAX_ROWS:
                    return out
                line = _line_from([_cell_text(value) for value in row])
                if line is None:
                    out.skipped += 1
                else:
                    out.lines.append(line)
    finally:
        workbook.close()
    return out


__all__ = ("MAX_IMPORT_BYTES", "Extracted", "extract_contacts", "looks_like_number")
